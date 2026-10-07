#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
科研进度汇报 —— 直接生成 .docx（无需复制粘贴、无需改后缀）

用法：
    python build_study_report.py 输入.json 输出.docx

输入 JSON 结构（除 student/studentId 外，字段缺失即跳过对应内容；不编造）：
{
  "student": "你的姓名",
  "studentId": "你的学号",
  "reportDate": "2026年10月8日",
  "period": "2026年10月1日 — 10月8日（共8天）",
  "platform": "中国大学MOOC、B站、arXiv",
  "summary": "深度学习基础与目标检测",
  "hours": "约18小时",
  "status": "按计划推进，节奏稳定",
  "literature":   [{"item": "《XXX》/ 作者 / 期刊", "content": "核心内容摘要"}],
  "theory":       [{"item": "课程名称", "content": "核心知识点（不写第X集）"}],
  "supplement":   [{"item": "知识模块", "content": "具体内容"}],
  "english":      [{"item": "内容", "content": "学习安排"}],
  "supplementIntro": "可选，知识补充板块的引导句；默认「围绕文献涉及的新知识，进行了系统学习与整理：」",
  "achievements": ["阶段性成果1", "阶段性成果2"],
  "plan": {"near": "近期计划内容", "next": "后续规划内容"}
}

其他可选字段：
  title            文档大标题，默认「科研进度汇报」
  compact          true 时收紧段间距与单元格内边距，把内容压进更少的页数（正文行距字号不变）
  contentSections  自定义「二、核心学习内容」下的章节；不写则用 literature/theory/supplement/english 四节
  supplementIntro  「知识补充」章节的引导句

contentSections 每一项：
  {"title": "章节名", "intro": "可选引导句", "type": "table" | "list",
   "columns": ["表头1", "表头2"], "ratios": [0.3, 0.7],     // type 缺省即 table
   "rows":    [{"item": "...", "content": "..."}],
   "items":   ["条目1", "条目2"]}                           // type=list 时用

样式与用户的 Word 模板（原 HTML/CSS 模板）一一对应，见 SKILL.md 的「样式映射」。
"""

import json
import sys

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, Twips

SONG = "宋体"
HEI = "黑体"
BODY_PT = 12.0        # 正文 12pt
TABLE_PT = 10.5       # 表格 10.5pt
H1_PT = 18.0
H2_PT = 14.0
H3_PT = 12.5
LINE = 1.8            # 正文行距
BORDER_SZ = 6         # 0.75pt ≈ 1px
BORDER_COLOR = "888888"
HEAD_FILL = "D9E1F2"
ACCENT = "2C7DA0"

# 紧凑排版：输入 JSON 里写 "compact": true 打开。只收紧段间距与单元格内边距，
# 不动正文 1.8 倍行距、字号和表格样式，用来把内容压进更少的页数。
COMPACT = False


# ---------------------------------------------------------------- 基础工具
def set_font(run, name=SONG, size=BODY_PT, bold=False, letter_spacing_pt=None):
    """设置字体，同时写 w:eastAsia，否则中文字形不会按指定字体渲染。"""
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)
    if letter_spacing_pt:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(int(letter_spacing_pt * 20)))
        sz = rPr.find(qn("w:sz"))          # w:spacing 必须排在 w:sz 之前
        if sz is not None:
            sz.addprevious(sp)
        else:
            rPr.append(sp)


def fmt_par(par, before=0, after=0, line=LINE, align=None, left_pt=None):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        par.alignment = align
    if left_pt is not None:
        pf.left_indent = Pt(left_pt)


def par_border(par, **edges):
    """edges: top/left/bottom/right = (sz_eighth_pt, color, space_pt)"""
    pPr = par._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.insert(0, pBdr)
    for side in ("top", "left", "bottom", "right"):
        if side not in edges:
            continue
        sz, color, space = edges[side]
        el = OxmlElement("w:" + side)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(int(sz)))
        el.set(qn("w:space"), str(int(space)))
        el.set(qn("w:color"), color)
        pBdr.append(el)


def table_borders(table, sz=BORDER_SZ, color=BORDER_COLOR):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def cell_margins(table, pt=None):
    if pt is None:
        pt = 3.0 if COMPACT else 5.0
    dxa = str(int(pt * 20))
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + side)
        el.set(qn("w:w"), dxa)
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def shade(cell, fill=HEAD_FILL):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_text(cell, text, bold=False, size=TABLE_PT, font=SONG, center=False):
    cell.text = ""
    par = cell.paragraphs[0]
    fmt_par(par, before=0, after=0, line=1.4, align=WD_ALIGN_PARAGRAPH.CENTER if center else None)
    run = par.add_run(str(text))
    set_font(run, font, size, bold)


def set_col_widths(table, ratios):
    """按比例固定列宽：同时改写 tblGrid、tblW 和单元格宽度，避免 Word 自行均分。"""
    total = 9639  # twips = 170mm，等于正文可用宽度
    table.autofit = False
    tblPr = table._tbl.tblPr

    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")

    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")

    widths = [int(total * r) for r in ratios]
    widths[-1] = total - sum(widths[:-1])

    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(widths):
                cell.width = Twips(widths[idx])


def fix_view_zoom(doc):
    """python-docx 默认模板带 <w:zoom w:val="bestFit"/>，会让部分渲染器按缩放视口出图
    （A4 被放大 2 倍后只截左上角）。固定为 100%，只影响打开时的视图比例。"""
    settings = doc.settings.element
    zoom = settings.find(qn("w:zoom"))
    if zoom is None:
        zoom = OxmlElement("w:zoom")
        settings.insert(0, zoom)
    if zoom.get(qn("w:val")) is not None:
        del zoom.attrib[qn("w:val")]
    zoom.set(qn("w:percent"), "100")


def header_row(row):
    """表头行：浅蓝底纹、加粗居中，并"与下段同页"——否则表头会孤零零留在页尾，
    再在下一页重复一次，看起来像多出来的空表头。"""
    for cell in row.cells:
        shade(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for par in cell.paragraphs:
            par.paragraph_format.keep_with_next = True


def table_rows_no_split(table, repeat_header=True):
    """禁止表格行跨页断开（否则表头会被切成两半），跨页时重复表头。"""
    for idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn("w:cantSplit")) is None:
            trPr.append(OxmlElement("w:cantSplit"))
        if idx == 0 and repeat_header and trPr.find(qn("w:tblHeader")) is None:
            trPr.append(OxmlElement("w:tblHeader"))


# ---------------------------------------------------------------- 文档结构
def add_title(doc, text):
    par = doc.add_paragraph()
    fmt_par(par, after=18, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = par.add_run(text)
    set_font(run, HEI, H1_PT, bold=True, letter_spacing_pt=2)


def add_h2(doc, text):
    par = doc.add_paragraph()
    fmt_par(par, before=12 if COMPACT else 16, after=6 if COMPACT else 8, line=1.2, left_pt=8)
    par_border(par, left=(24, ACCENT, 8))  # 3pt 左边框 + 8pt 间距
    run = par.add_run(text)
    set_font(run, HEI, H2_PT, bold=True)


def add_h3(doc, text):
    par = doc.add_paragraph()
    fmt_par(par, before=8 if COMPACT else 12, after=3 if COMPACT else 4, line=1.2)
    run = par.add_run(text)
    set_font(run, HEI, H3_PT, bold=True)


def add_body(doc, text):
    par = doc.add_paragraph()
    fmt_par(par)
    run = par.add_run(text)
    set_font(run, SONG, BODY_PT)
    return par


def add_bullet(doc, text, label=None):
    par = doc.add_paragraph(style="List Bullet")
    pad = 1 if COMPACT else 2
    fmt_par(par, before=pad, after=pad, line=1.6, left_pt=20)
    if label:
        r1 = par.add_run(label)
        set_font(r1, SONG, BODY_PT, bold=True)
    r2 = par.add_run(text)
    set_font(r2, SONG, BODY_PT)
    return par


def add_two_col_table(doc, header, rows, ratios):
    table = doc.add_table(rows=1, cols=2)
    table_borders(table)
    cell_margins(table)
    hdr = table.rows[0]
    cell_text(hdr.cells[0], header[0], bold=True, center=True)
    cell_text(hdr.cells[1], header[1], bold=True, center=True)
    header_row(hdr)
    for row in rows:
        cells = table.add_row().cells
        cell_text(cells[0], row.get("item", ""))
        cell_text(cells[1], row.get("content", ""))
    set_col_widths(table, ratios)
    table_rows_no_split(table)
    return table


def build(data, out_path):
    global COMPACT
    COMPACT = bool(data.get("compact"))
    doc = Document()

    # 页面：A4 + 上下 25mm、左右 20mm
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = Mm(25)
    sec.left_margin = sec.right_margin = Mm(20)

    # 默认字体
    normal = doc.styles["Normal"]
    normal.font.name = SONG
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)

    add_title(doc, data.get("title", "科研进度汇报"))

    # 一、基本信息
    add_h2(doc, "一、基本信息")
    info = [
        ("汇报人：", data.get("student", "")),
        ("学号：", data.get("studentId", "")),
        ("汇报日期：", data.get("reportDate", "")),
        ("学习周期：", data.get("period", "")),
        ("学习平台：", data.get("platform", "")),
        ("学习内容：", data.get("summary", "")),
        ("累计用时：", data.get("hours", "")),
        ("学习状态：", data.get("status", "")),
    ]
    for label, value in info:
        if value:
            add_bullet(doc, value, label)

    # 二、核心学习内容 —— 章节可用 contentSections 完全替换；不写就用下面四节的固定顺序
    add_h2(doc, "二、核心学习内容")

    # 小标题编号按"实际出现的章节"重排：缺了文献阅读，理论学习就是（一）
    sub = ["（一）", "（二）", "（三）", "（四）", "（五）", "（六）", "（七）", "（八）"]
    counter = {"n": 0}

    def sub_title(name):
        counter["n"] += 1
        return sub[counter["n"] - 1] + name

    def add_section(sec):
        """渲染一个章节：标题 + 可选引导句 + 二列表格（默认）或项目符号列表。"""
        rows = sec.get("rows") or []
        items = sec.get("items") or []
        if not rows and not items:
            return
        add_h3(doc, sub_title(sec.get("title", "学习内容")))
        if sec.get("intro"):
            add_body(doc, sec["intro"])
        if sec.get("type") == "list" or (items and not rows):
            for item in items:
                add_bullet(doc, item)
        else:
            columns = sec.get("columns") or ["项目", "内容"]
            ratios = sec.get("ratios") or [0.22, 0.78]
            add_two_col_table(doc, tuple(columns[:2]), rows, tuple(ratios[:2]))

    sections = data.get("contentSections")
    if sections is None:
        sections = [
            {"title": "文献阅读", "columns": ["文献", "核心内容"], "ratios": [0.32, 0.68],
             "rows": data.get("literature") or []},
            {"title": "理论学习", "columns": ["课程", "核心知识点"], "ratios": [0.22, 0.78],
             "rows": data.get("theory") or []},
            {"title": "知识补充", "columns": ["知识模块", "具体内容"], "ratios": [0.22, 0.78],
             "intro": data.get("supplementIntro", "围绕文献涉及的新知识，进行了系统学习与整理："),
             "rows": data.get("supplement") or []},
            {"title": "英语学习", "columns": ["内容", "学习安排"], "ratios": [0.22, 0.78],
             "rows": data.get("english") or []},
        ]
    for sec in sections:
        add_section(sec)

    # 三、阶段性学习成果
    if data.get("achievements"):
        add_h2(doc, "三、阶段性学习成果")
        for item in data["achievements"]:
            add_bullet(doc, item)

    # 四、后续学习规划
    plan = data.get("plan") or {}
    if plan.get("near") or plan.get("next"):
        add_h2(doc, "四、后续学习规划")
        add_body(doc, "根据导师指导，整体学习路径规划如下：")
        table = doc.add_table(rows=1, cols=3)
        table_borders(table)
        cell_margins(table)
        hdr = table.rows[0]
        for i, name in enumerate(("阶段", "时间", "计划内容")):
            cell_text(hdr.cells[i], name, bold=True, center=True)
        header_row(hdr)
        for stage, when, content in (
            ("近期计划", "接下来", plan.get("near", "")),
            ("后续规划", "下一阶段", plan.get("next", "")),
        ):
            cells = table.add_row().cells
            cell_text(cells[0], stage, bold=True)
            cell_text(cells[1], when)
            cell_text(cells[2], content)
        set_col_widths(table, (0.18, 0.18, 0.64))
        table_rows_no_split(table)

    # 落款
    sign = doc.add_paragraph()
    fmt_par(sign, before=12 if COMPACT else 20, after=0, line=1.2, align=WD_ALIGN_PARAGRAPH.RIGHT)
    par_border(sign, top=(BORDER_SZ, "CCCCCC", 8 if COMPACT else 12))
    _stu = data.get("student", "")
    _sign = ("学生：%s　　" % _stu if _stu else "") + "日期：%s" % data.get("reportDate", "")
    run = sign.add_run(_sign)
    set_font(run, SONG, BODY_PT)

    fix_view_zoom(doc)
    doc.save(out_path)
    return out_path


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return 2
    with open(sys.argv[1], "r", encoding="utf-8") as fh:
        data = json.load(fh)
    path = build(data, sys.argv[2])
    print("已生成：" + path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
